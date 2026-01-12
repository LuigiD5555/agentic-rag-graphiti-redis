"""Module for docx files loader."""
from typing import List
from zipfile import BadZipFile

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from langchain_core.documents import Document

from src.workflows.ingestion.loaders.errors import (
    LoaderInvalidFormatError,
    ensure_file_exists,
)
from src.workflows.ingestion.loaders.office_client import OfficeToolClient


class DocxLoader:
    """Docx document loader."""

    def __init__(self, path: str):
        self._path = path
        self._office_client = OfficeToolClient()

    def load(self) -> List[Document]:
        """
        Load docx documents.

        Returns:
            List[Document]: Loaded documents.
        """
        ensure_file_exists(self._path)
        return self._python_docx_loader()

    def _python_docx_loader(self) -> List[Document]:
        """Loader that relies on python-docx."""
        try:
            doc = DocxDocument(self._path)
        except (PackageNotFoundError, BadZipFile) as exc:
            return self._unstructured_loader(exc)

        lines: List[str] = []

        def _append(text: str) -> None:
            stripped = text.strip()
            if stripped:
                lines.append(stripped)

        for paragraph in doc.paragraphs:
            _append(paragraph.text)

        for table in doc.tables:
            for row in table.rows:
                row_values = []
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        row_values.append(cell_text)
                _append(" | ".join(row_values))

        return [
            Document(
                page_content="\n".join(lines),
                metadata={"source": self._path},
            )
        ]

    def _unstructured_loader(self, original_exc: Exception) -> List[Document]:
        """Fallback loader via rag-tool-document-processor for oddball docx files."""
        try:
            doc = self._office_client.load_as_document(self._path)
            return [doc]
        except RuntimeError as exc:
            plain = self._plain_text_fallback(exc)
            if plain:
                return plain
            raise LoaderInvalidFormatError(
                self._path,
                expected="DOCX",
                detail=str(exc),
            ) from original_exc

    def _plain_text_fallback(self, secondary_exc: Exception) -> List[Document]:
        """
        Last-resort fallback: attempt to read file bytes as text.

        Useful when a file is mislabeled as .docx but not a valid Zip/docx. We
        return an empty list if the content is unreadable to avoid crashing the
        pipeline.
        """
        try:
            with open(self._path, "rb") as handle:
                raw = handle.read()
        except OSError:
            return []

        # Try UTF-8 first, then fallback to other encodings
        text = None
        for encoding in ("utf-8", "utf-16", "cp1252", "latin-1"):
            try:
                text = raw.decode(encoding).strip()
                if text:
                    break
            except (UnicodeDecodeError, LookupError):
                continue

        # Final fallback: decode with errors='ignore'
        if not text:
            text = raw.decode("utf-8", errors="ignore").strip()

        if not text:
            return []

        # Clean any surrogate characters
        text = text.encode('utf-8', errors='surrogatepass').decode('utf-8', errors='ignore')

        return [
            Document(
                page_content=text,
                metadata={
                    "source": self._path,
                    "fallback": "plain_text",
                    "error": str(secondary_exc),
                },
            )
        ]
